"""Génère le classeur modèle de codification des emplacements (CAMEG Atakpamé).

Usage : python outils/generer_modele.py [chemin_sortie.xlsx]
"""
import sys
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SORTIE = sys.argv[1] if len(sys.argv) > 1 else "Modele_Codification_Stock_CAMEG.xlsx"

N_EMPL = 1000    # lignes prévues dans Emplacements
N_PROD = 1000    # lignes prévues dans Produits
N_STOCK = 2000   # lignes prévues dans Stock
N_RECH = 40      # résultats affichés dans Recherche

POLICE = "Arial"
F_NORMAL = Font(name=POLICE, size=10)
F_TITRE = Font(name=POLICE, size=16, bold=True, color="1F4E79")
F_SOUS_TITRE = Font(name=POLICE, size=12, bold=True, color="1F4E79")
F_ENTETE = Font(name=POLICE, size=10, bold=True, color="FFFFFF")
F_EXEMPLE = Font(name=POLICE, size=10, color="0000FF")
F_FORMULE = Font(name=POLICE, size=10, color="000000")
FOND_ENTETE = PatternFill("solid", fgColor="1F4E79")
FOND_SAISIE = PatternFill("solid", fgColor="FFF2CC")
FOND_FORMULE = PatternFill("solid", fgColor="EDEDED")
FOND_ALERTE = PatternFill("solid", fgColor="F8CBAD")
FOND_ATTENTION = PatternFill("solid", fgColor="FFE699")
FOND_OK = PatternFill("solid", fgColor="C6EFCE")
TRAIT = Side(style="thin", color="BFBFBF")
BORDURE = Border(left=TRAIT, right=TRAIT, top=TRAIT, bottom=TRAIT)
CENTRE = Alignment(horizontal="center", vertical="center", wrap_text=True)
GAUCHE = Alignment(horizontal="left", vertical="top", wrap_text=True)


def entetes(ws, titres, largeurs, ligne=1):
    for col, (titre, largeur) in enumerate(zip(titres, largeurs), start=1):
        c = ws.cell(row=ligne, column=col, value=titre)
        c.font, c.fill, c.alignment, c.border = F_ENTETE, FOND_ENTETE, CENTRE, BORDURE
        ws.column_dimensions[c.column_letter].width = largeur
    ws.row_dimensions[ligne].height = 32
    ws.freeze_panes = ws.cell(row=ligne + 1, column=1)


def zone_saisie(ws, colonnes, premiere, derniere):
    for col in colonnes:
        for r in range(premiere, derniere + 1):
            c = ws.cell(row=r, column=col)
            c.fill, c.border, c.font = FOND_SAISIE, BORDURE, F_NORMAL


def zone_formule(ws, col, premiere, derniere, modele):
    for r in range(premiere, derniere + 1):
        c = ws.cell(row=r, column=col, value=modele.format(r=r, p=r - 1))
        c.fill, c.border, c.font = FOND_FORMULE, BORDURE, F_FORMULE


def exemple(ws, ligne, valeurs, col_debut=1):
    for i, v in enumerate(valeurs):
        if v is None:
            continue
        c = ws.cell(row=ligne, column=col_debut + i, value=v)
        c.font = F_EXEMPLE


wb = Workbook()

# ---------------------------------------------------------------- Lisez-moi
lm = wb.active
lm.title = "Lisez-moi"
lm.sheet_view.showGridLines = False
lm.column_dimensions["A"].width = 3
lm.column_dimensions["B"].width = 28
lm.column_dimensions["C"].width = 95
lignes = [
    ("t", "Système de codification des emplacements – Magasin CAMEG Atakpamé"),
    ("p", "Modèle (squelette) à adapter. Objectif : que chaque produit ait une adresse unique et que n'importe qui puisse le retrouver."),
    ("", ""),
    ("s", "1. Format du code d'emplacement"),
    ("kv", ("Format", "ZONE + ÉTAGÈRE  -  NIVEAU  -  POSITION        exemple :  A03-2-04")),
    ("kv", ("Zone (1 lettre)", "Type de stockage : A = général, F = froid +2/+8 °C, S = sécurisé, N = inflammables, P = palettes/vrac, Q = quarantaine, R = réception, E = expédition (voir onglet Zones)")),
    ("kv", ("Étagère (2 chiffres)", "Numéro de l'étagère / rack dans la zone : 01, 02 … 99. Reprendre si possible la numérotation déjà peinte sur les étagères (colonne « Ancien code »).")),
    ("kv", ("Niveau (1 chiffre)", "1 = niveau le plus bas (sol), puis 2, 3, 4 … en montant.")),
    ("kv", ("Position (2 chiffres)", "Case / casier de gauche à droite quand on fait face à l'étagère : 01, 02, 03 …")),
    ("kv", ("Lecture", "A03-2-04 = zone stock général, étagère 03, 2e niveau en partant du bas, 4e case depuis la gauche.")),
    ("", ""),
    ("s", "2. Règles d'or"),
    ("kv", ("Un code = un endroit", "Jamais deux emplacements avec le même code ; on ne réutilise pas un code supprimé pour un autre endroit.")),
    ("kv", ("Longueur fixe", "Toujours 01, 02… (avec le zéro) pour que le tri alphabétique suive l'ordre physique dans le magasin.")),
    ("kv", ("Code produit ≠ code emplacement", "Le code article ne contient PAS l'emplacement. Le lien se fait dans les onglets Produits (emplacement principal) et Stock (où se trouve chaque lot).")),
    ("kv", ("Étiqueter physiquement", "Chaque case doit porter son code (onglet Étiquettes à imprimer). Un code qui n'est pas affiché sur l'étagère ne sert à rien.")),
    ("kv", ("FEFO", "Premier périmé, premier sorti : sur un même emplacement, le lot qui périme le plus tôt devant. L'onglet Stock signale les lots proches de la péremption.")),
    ("kv", ("Mettre à jour", "Tout déplacement de produit = mise à jour de l'onglet Stock le jour même, sinon le système ment.")),
    ("", ""),
    ("s", "3. Comment remplir ce classeur"),
    ("kv", ("Ordre", "① Zones (vérifier)  →  ② Emplacements (inventaire des étagères)  →  ③ Produits  →  ④ Stock (où est chaque lot)  →  ⑤ Recherche / Étiquettes")),
    ("kv", ("Cellules jaunes", "À remplir / modifier par vous.")),
    ("kv", ("Cellules grises", "Formules automatiques : ne pas taper dedans.")),
    ("kv", ("Texte bleu", "Données d'EXEMPLE (codes EX-…, emplacements fictifs) : à effacer et remplacer par les vraies données du magasin.")),
    ("kv", ("Recherche", "Onglet Recherche : taper une partie du nom (ex. « amox ») ou un code dans la case jaune → la liste des emplacements s'affiche.")),
    ("kv", ("Codes articles", "Utiliser les codes articles officiels de la CAMEG s'ils existent (catalogue / logiciel de gestion). Ne pas en inventer de nouveaux si un code existe déjà.")),
    ("", ""),
    ("s", "4. Références utilisées"),
    ("p", "• USAID | DELIVER PROJECT – The Logistics Handbook (2011) et Guidelines for the Storage of Essential Medicines and Other Health Commodities (JSI/USAID, 2003)."),
    ("p", "• OMS – Bonnes pratiques de stockage et de distribution des produits médicaux (WHO TRS n° 1025, annexe 7, 2020)."),
    ("p", "• MSH – MDS-3 : Managing Access to Medicines and Health Technologies (2012), chapitres sur la gestion des magasins."),
    ("p", "• GS1 – identification des lieux (GLN) et codes-barres des médicaments, pour une évolution future vers le scan."),
    ("p", "Note : il n'existe pas de norme internationale imposant un format de code d'emplacement ; le format Zone-Allée/Étagère-Niveau-Position est la convention utilisée par la plupart des entrepôts et logiciels (WMS)."),
    ("", ""),
    ("p", f"Modèle généré le {date.today():%d/%m/%Y}."),
]
r = 1
for kind, val in lignes:
    if kind == "t":
        lm.cell(row=r, column=2, value=val).font = F_TITRE
        lm.row_dimensions[r].height = 26
    elif kind == "s":
        lm.cell(row=r, column=2, value=val).font = F_SOUS_TITRE
        lm.row_dimensions[r].height = 20
    elif kind == "p":
        lm.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        c = lm.cell(row=r, column=2, value=val)
        c.font, c.alignment = F_NORMAL, GAUCHE
        lm.row_dimensions[r].height = 15 if len(val) < 120 else 28
    elif kind == "kv":
        k, v = val
        a = lm.cell(row=r, column=2, value=k)
        a.font, a.alignment, a.border = Font(name=POLICE, size=10, bold=True), GAUCHE, BORDURE
        b = lm.cell(row=r, column=3, value=v)
        b.font, b.alignment, b.border = F_NORMAL, GAUCHE, BORDURE
        lm.row_dimensions[r].height = 15 if len(v) < 95 else 28
    r += 1
lm.page_setup.orientation = "landscape"
lm.page_setup.fitToWidth = 1
lm.page_setup.fitToHeight = 0
lm.sheet_properties.pageSetUpPr.fitToPage = True
lm["C5"].font = Font(name="Consolas", size=11, bold=True)
lm["C22"].fill = FOND_SAISIE
lm["C23"].fill = FOND_FORMULE
lm["C24"].font = F_EXEMPLE

# ---------------------------------------------------------------- Zones
zo = wb.create_sheet("Zones")
entetes(zo, ["Code zone", "Libellé", "Conditions de stockage", "Exemples de produits"], [11, 34, 48, 48])
zones = [
    ("A", "Stock général", "Température ambiante contrôlée (selon étiquette, en général < 25 °C ou < 30 °C), au sec, à l'abri de la lumière, hors sol", "Comprimés, gélules, sirops, pommades"),
    ("F", "Chaîne de froid", "+2 °C à +8 °C, relevé de température 2 fois/jour", "Ocytocine, insulines, certains vaccins et réactifs"),
    ("S", "Sécurisé", "Armoire ou local fermé à clé, accès limité, registre", "Stupéfiants, psychotropes, produits de grande valeur"),
    ("N", "Inflammables", "Local séparé et ventilé, loin des sources de chaleur", "Alcool, éther, produits à base de solvants"),
    ("P", "Palettes / vrac", "Sur palettes, jamais à même le sol, hauteur de gerbage respectée", "Perfusions, SRO en cartons, gros volumes"),
    ("Q", "Quarantaine", "Séparée et signalée (« NE PAS DISTRIBUER »)", "Périmés, endommagés, retours, lots rappelés ou en attente de contrôle"),
    ("R", "Réception", "Zone tampon à l'entrée, contrôle avant rangement", "Livraisons en cours de vérification"),
    ("E", "Expédition", "Zone de préparation des commandes", "Commandes préparées en attente d'enlèvement"),
]
for i, ligne in enumerate(zones, start=2):
    for j, v in enumerate(ligne, start=1):
        c = zo.cell(row=i, column=j, value=v)
        c.font, c.fill, c.border, c.alignment = F_NORMAL, FOND_SAISIE, BORDURE, GAUCHE
    zo.cell(row=i, column=1).alignment = CENTRE
    zo.cell(row=i, column=1).font = Font(name=POLICE, size=11, bold=True)
    zo.row_dimensions[i].height = 30
for i in range(len(zones) + 2, 16):
    for j in range(1, 5):
        c = zo.cell(row=i, column=j)
        c.fill, c.border, c.font = FOND_SAISIE, BORDURE, F_NORMAL
zo.cell(row=17, column=1, value="Lettres I et O volontairement évitées (confusion avec 1 et 0). Ajouter une zone si besoin dans les lignes jaunes vides.").font = Font(name=POLICE, size=9, italic=True)
ZONES_REF = "Zones!$A$2:$A$15"

# ---------------------------------------------------------------- Emplacements
em = wb.create_sheet("Emplacements")
entetes(em, ["Code emplacement", "Zone", "N° étagère", "Niveau\n(1 = bas)", "Position\n(gauche→droite)",
             "Libellé zone", "Ancien code\n(peint sur l'étagère)", "Type de support", "Statut", "Remarque"],
        [18, 8, 11, 11, 15, 20, 18, 18, 11, 30])
em["A1"].comment = Comment("Calculé automatiquement : Zone + Étagère(2 chiffres) - Niveau - Position(2 chiffres).", "Modèle")
zone_saisie(em, [2, 3, 4, 5, 7, 8, 9, 10], 2, N_EMPL + 1)
zone_formule(em, 1, 2, N_EMPL + 1,
             '=IF(OR(B{r}="",C{r}="",D{r}="",E{r}=""),"",UPPER(B{r})&TEXT(C{r},"00")&"-"&D{r}&"-"&TEXT(E{r},"00"))')
zone_formule(em, 6, 2, N_EMPL + 1,
             '=IF(B{r}="","",IFERROR(INDEX(Zones!$B$2:$B$15,MATCH(B{r},Zones!$A$2:$A$15,0)),"⚠ zone inconnue"))')
for col in (1, 6):
    for rr in range(2, N_EMPL + 2):
        em.cell(row=rr, column=col).alignment = CENTRE if col == 1 else GAUCHE
        if col == 1:
            em.cell(row=rr, column=col).font = Font(name="Consolas", size=10, bold=True)

ex_empl = []
for et, anc in ((1, "E1"), (2, "E2")):
    for niv in range(1, 5):
        for pos in range(1, 4):
            ex_empl.append(("A", et, niv, pos, anc, "Étagère métallique"))
ex_empl += [("F", 1, n, 1, "Frigo 1", "Réfrigérateur") for n in range(1, 4)]
ex_empl += [("S", 1, n, 1, "Armoire", "Armoire à clé") for n in range(1, 3)]
ex_empl += [("N", 1, 1, 1, "", "Étagère local ventilé")]
ex_empl += [("P", 1, 1, p, "", "Palette") for p in range(1, 5)]
ex_empl += [("Q", 1, 1, p, "", "Étagère") for p in range(1, 3)]
ex_empl += [("R", 1, 1, 1, "", "Sol (zone tampon)"), ("E", 1, 1, 1, "", "Table de préparation")]
for i, (z, et, niv, pos, anc, sup) in enumerate(ex_empl, start=2):
    exemple(em, i, [z, et, niv, pos], col_debut=2)
    exemple(em, i, [anc or None, sup, "Actif"], col_debut=7)

dv_zone = DataValidation(type="list", formula1=ZONES_REF, allow_blank=True,
                         error="Choisir une zone de l'onglet Zones", errorTitle="Zone")
dv_etag = DataValidation(type="whole", operator="between", formula1="1", formula2="99", allow_blank=True,
                         error="Nombre entier de 1 à 99", errorTitle="Étagère")
dv_niv = DataValidation(type="whole", operator="between", formula1="1", formula2="9", allow_blank=True,
                        error="Nombre entier de 1 à 9 (1 = niveau du bas)", errorTitle="Niveau")
dv_pos = DataValidation(type="whole", operator="between", formula1="1", formula2="99", allow_blank=True,
                        error="Nombre entier de 1 à 99", errorTitle="Position")
dv_statut = DataValidation(type="list", formula1='"Actif,Bloqué,Supprimé"', allow_blank=True)
for dv, col in ((dv_zone, "B"), (dv_etag, "C"), (dv_niv, "D"), (dv_pos, "E"), (dv_statut, "I")):
    em.add_data_validation(dv)
    dv.add(f"{col}2:{col}{N_EMPL + 1}")
# doublons de code
em.conditional_formatting.add(f"A2:A{N_EMPL + 1}", FormulaRule(
    formula=[f'AND(A2<>"",COUNTIF($A$2:$A${N_EMPL + 1},A2)>1)'], fill=FOND_ALERTE))
em.auto_filter.ref = f"A1:J{N_EMPL + 1}"

# ---------------------------------------------------------------- Produits
pr = wb.create_sheet("Produits")
entetes(pr, ["Code article", "DCI (nom générique)", "Dosage", "Forme", "Conditionnement",
             "Désignation complète", "Classe thérapeutique", "Zone requise",
             "Emplacement principal", "Contrôle emplacement", "Classe ABC", "VEN"],
        [13, 28, 16, 16, 20, 42, 22, 10, 16, 22, 9, 7])
pr["A1"].comment = Comment("Mettre le code article officiel CAMEG. Les codes EX-… sont des exemples.", "Modèle")
pr["H1"].comment = Comment("Lettre de zone où le produit DOIT être stocké (F pour la chaîne de froid, etc.).", "Modèle")
pr["I1"].comment = Comment("Emplacement fixe habituel du produit (casier de prélèvement). Les autres lots peuvent être ailleurs : voir onglet Stock.", "Modèle")
pr["K1"].comment = Comment("A = produits qui représentent ~80 % de la valeur consommée, B = ~15 %, C = ~5 %. Mettre les A près de la sortie.", "Modèle")
pr["L1"].comment = Comment("V = vital, E = essentiel, N = non essentiel (classification OMS/MSH).", "Modèle")
zone_saisie(pr, [1, 2, 3, 4, 5, 7, 8, 9, 11, 12], 2, N_PROD + 1)
zone_formule(pr, 6, 2, N_PROD + 1, '=IF(B{r}="","",TRIM(B{r}&" "&C{r}&" "&D{r}))')
zone_formule(pr, 10, 2, N_PROD + 1,
             '=IF(I{r}="","",IF(COUNTIF(Emplacements!$A$2:$A$' + str(N_EMPL + 1) + ',I{r})=0,"⚠ emplacement inconnu",'
             'IF(AND(H{r}<>"",LEFT(I{r},1)<>H{r}),"⚠ mauvaise zone","OK")))')
ex_prod = [
    ("EX-0001", "Paracétamol", "500 mg", "comprimé", "boîte de 1000", "Antalgique", "A", "A01-1-01", "A", "E"),
    ("EX-0002", "Amoxicilline", "500 mg", "gélule", "boîte de 1000", "Anti-infectieux", "A", "A01-1-02", "A", "V"),
    ("EX-0003", "Artéméther + Luméfantrine", "20/120 mg", "comprimé", "boîte de 30 x 24 cp", "Antipaludique", "A", "A01-1-03", "A", "V"),
    ("EX-0004", "Fer + Acide folique", "200/0,4 mg", "comprimé", "boîte de 1000", "Antianémique", "A", "A02-1-01", "B", "E"),
    ("EX-0005", "Gants d'examen", "taille M", "non stériles", "boîte de 100", "Consommable", "A", "A02-2-01", "C", "E"),
    ("EX-0006", "Sels de réhydratation orale", "20,5 g", "sachet", "carton de 100", "Réhydratation", "P", "P01-1-01", "B", "V"),
    ("EX-0007", "Chlorure de sodium", "0,9 % 500 ml", "perfusion", "carton de 20", "Soluté", "P", "P01-1-02", "B", "V"),
    ("EX-0008", "Ocytocine", "10 UI/ml", "injectable", "boîte de 100 amp.", "Ocytocique", "F", "F01-1-01", "B", "V"),
    ("EX-0009", "Diazépam", "10 mg/2 ml", "injectable", "boîte de 100 amp.", "Psychotrope", "S", "S01-1-01", "C", "V"),
    ("EX-0010", "Alcool éthylique", "70 %", "solution", "flacon de 1 L", "Antiseptique", "N", "N01-1-01", "C", "E"),
]
for i, row in enumerate(ex_prod, start=2):
    exemple(pr, i, row[:5], 1)
    exemple(pr, i, row[5:9], 7)
    exemple(pr, i, row[9:], 12)
dv_zone_p = DataValidation(type="list", formula1=ZONES_REF, allow_blank=True)
dv_abc = DataValidation(type="list", formula1='"A,B,C"', allow_blank=True)
dv_ven = DataValidation(type="list", formula1='"V,E,N"', allow_blank=True)
for dv, col in ((dv_zone_p, "H"), (dv_abc, "K"), (dv_ven, "L")):
    pr.add_data_validation(dv)
    dv.add(f"{col}2:{col}{N_PROD + 1}")
pr.conditional_formatting.add(f"J2:J{N_PROD + 1}", FormulaRule(formula=['LEFT(J2,1)="⚠"'], fill=FOND_ALERTE))
pr.conditional_formatting.add(f"J2:J{N_PROD + 1}", FormulaRule(formula=['J2="OK"'], fill=FOND_OK))
pr.conditional_formatting.add(f"A2:A{N_PROD + 1}", FormulaRule(
    formula=[f'AND(A2<>"",COUNTIF($A$2:$A${N_PROD + 1},A2)>1)'], fill=FOND_ALERTE))
pr.auto_filter.ref = f"A1:L{N_PROD + 1}"

# ---------------------------------------------------------------- Stock
st = wb.create_sheet("Stock")
entetes(st, ["Emplacement", "Code article", "Désignation", "N° de lot", "Date de péremption",
             "Quantité\n(unités)", "Date dernier\ncomptage", "Mois avant\npéremption", "Alerte",
             "Contrôle emplacement", "(technique)\nn° recherche"],
        [14, 13, 40, 14, 14, 12, 14, 12, 30, 22, 12])
st["A1"].comment = Comment("Une ligne = un lot à un emplacement. Si un lot est réparti sur 2 emplacements → 2 lignes.", "Modèle")
st["K1"].comment = Comment("Colonne technique utilisée par l'onglet Recherche. Ne pas modifier.", "Modèle")
zone_saisie(st, [1, 2, 4, 5, 6, 7], 2, N_STOCK + 1)
PROD_A = f"Produits!$A$2:$A${N_PROD + 1}"
PROD_F = f"Produits!$F$2:$F${N_PROD + 1}"
zone_formule(st, 3, 2, N_STOCK + 1,
             f'=IF(B{{r}}="","",IFERROR(INDEX({PROD_F},MATCH(B{{r}},{PROD_A},0)),"⚠ code article inconnu"))')
zone_formule(st, 8, 2, N_STOCK + 1, '=IF(E{r}="","",ROUND((E{r}-TODAY())/30.4,1))')
zone_formule(st, 9, 2, N_STOCK + 1,
             '=IF(E{r}="","",IF(H{r}<0,IF(LEFT(A{r},1)="Q","Périmé (en quarantaine)","PÉRIMÉ → déplacer en zone Q"),'
             'IF(H{r}<3,"Péremption < 3 mois",IF(H{r}<6,"Péremption < 6 mois",""))))')
zone_formule(st, 10, 2, N_STOCK + 1,
             '=IF(A{r}="","",IF(COUNTIF(Emplacements!$A$2:$A$' + str(N_EMPL + 1) + ',A{r})=0,"⚠ emplacement inconnu","OK"))')
zone_formule(st, 11, 2, N_STOCK + 1,
             '=IF(AND(B{r}<>"",ISNUMBER(SEARCH(Recherche!$C$3,C{r}&" "&B{r}&" "&A{r}&" "&D{r}))),MAX($K$1:K{p})+1,"")')
for rr in range(2, N_STOCK + 2):
    st.cell(row=rr, column=5).number_format = "DD/MM/YYYY"
    st.cell(row=rr, column=7).number_format = "DD/MM/YYYY"
    st.cell(row=rr, column=6).number_format = "#,##0"
    st.cell(row=rr, column=8).number_format = "0.0"
    st.cell(row=rr, column=11).font = Font(name=POLICE, size=8, color="808080")
ex_stock = [
    ("A01-1-01", "EX-0001", "PCM2405", date(2027, 11, 30), 12000, date(2026, 10, 1)),
    ("A01-1-01", "EX-0001", "PCM2502", date(2028, 4, 30), 20000, date(2026, 10, 1)),
    ("A01-1-02", "EX-0002", "AMX2311", date(2026, 12, 31), 8000, date(2026, 10, 1)),
    ("A01-1-03", "EX-0003", "ALU2408", date(2027, 6, 30), 450, date(2026, 10, 1)),
    ("A02-1-01", "EX-0004", "FAF2410", date(2027, 9, 30), 15000, date(2026, 10, 1)),
    ("A02-2-01", "EX-0005", "GTM2501", date(2029, 1, 31), 300, date(2026, 10, 1)),
    ("P01-1-01", "EX-0006", "SRO2406", date(2027, 5, 31), 2400, date(2026, 10, 1)),
    ("P01-1-02", "EX-0007", "NAC2409", date(2027, 8, 31), 640, date(2026, 10, 1)),
    ("F01-1-01", "EX-0008", "OXY2503", date(2027, 2, 28), 1500, date(2026, 10, 1)),
    ("S01-1-01", "EX-0009", "DZP2402", date(2027, 1, 31), 200, date(2026, 10, 1)),
    ("N01-1-01", "EX-0010", "ALC2501", date(2028, 1, 31), 60, date(2026, 10, 1)),
    ("Q01-1-01", "EX-0002", "AMX2209", date(2026, 8, 31), 500, date(2026, 10, 1)),
]
for i, row in enumerate(ex_stock, start=2):
    exemple(st, i, row[:2], 1)
    exemple(st, i, row[2:], 4)
dv_empl = DataValidation(type="list", formula1=f"Emplacements!$A$2:$A${N_EMPL + 1}", allow_blank=True,
                         showErrorMessage=False)
dv_art = DataValidation(type="list", formula1=PROD_A, allow_blank=True, showErrorMessage=False)
st.add_data_validation(dv_empl)
st.add_data_validation(dv_art)
dv_empl.add(f"A2:A{N_STOCK + 1}")
dv_art.add(f"B2:B{N_STOCK + 1}")
# FIND est sensible à la casse (contrairement à « = » dans Excel)
st.conditional_formatting.add(f"I2:I{N_STOCK + 1}", FormulaRule(formula=['ISNUMBER(FIND("PÉRIMÉ",I2))'], fill=FOND_ALERTE))
st.conditional_formatting.add(f"I2:I{N_STOCK + 1}", FormulaRule(formula=['ISNUMBER(FIND("quarantaine",I2))'], fill=FOND_FORMULE))
st.conditional_formatting.add(f"I2:I{N_STOCK + 1}", FormulaRule(formula=['ISNUMBER(FIND("Péremption",I2))'], fill=FOND_ATTENTION))
st.conditional_formatting.add(f"J2:J{N_STOCK + 1}", FormulaRule(formula=['LEFT(J2,1)="⚠"'], fill=FOND_ALERTE))
st.conditional_formatting.add(f"C2:C{N_STOCK + 1}", FormulaRule(formula=['LEFT(C2,1)="⚠"'], fill=FOND_ALERTE))
st.auto_filter.ref = f"A1:K{N_STOCK + 1}"

# ---------------------------------------------------------------- Recherche
re_ = wb.create_sheet("Recherche")
re_.sheet_view.showGridLines = False
re_["A1"] = "Où se trouve le produit ?"
re_["A1"].font = F_TITRE
re_["A3"] = "Tapez un nom, un code ou un lot :"
re_["A3"].font = Font(name=POLICE, size=11, bold=True)
re_.merge_cells("A3:B3")
re_["C3"] = "amox"
re_["C3"].font = Font(name=POLICE, size=14, bold=True, color="0000FF")
re_["C3"].fill = FOND_SAISIE
re_["C3"].border = Border(left=Side("medium"), right=Side("medium"), top=Side("medium"), bottom=Side("medium"))
re_["D3"] = f'=COUNT(Stock!$K$2:$K${N_STOCK + 1})&" ligne(s) trouvée(s)"'
re_["D3"].font = Font(name=POLICE, size=10, italic=True)
re_["A4"] = "Laisser vide pour tout afficher. Les 40 premiers résultats sont affichés ; utiliser le filtre de l'onglet Stock pour plus."
re_["A4"].font = Font(name=POLICE, size=9, italic=True, color="808080")
cols_rech = [("N°", None, 5), ("Emplacement", "A", 15), ("Désignation", "C", 42), ("Code article", "B", 13),
             ("N° de lot", "D", 13), ("Péremption", "E", 13), ("Quantité", "F", 11), ("Alerte", "I", 30)]
for j, (titre, _, larg) in enumerate(cols_rech, start=1):
    c = re_.cell(row=6, column=j, value=titre)
    c.font, c.fill, c.alignment, c.border = F_ENTETE, FOND_ENTETE, CENTRE, BORDURE
    re_.column_dimensions[c.column_letter].width = larg
K = f"Stock!$K$2:$K${N_STOCK + 1}"
for n in range(1, N_RECH + 1):
    rr = 6 + n
    c = re_.cell(row=rr, column=1, value=n)
    c.font, c.alignment, c.border = Font(name=POLICE, size=9, color="808080"), CENTRE, BORDURE
    for j, (_, src, _) in enumerate(cols_rech[1:], start=2):
        f = f'=IFERROR(INDEX(Stock!${src}$2:${src}${N_STOCK + 1},MATCH($A{rr},{K},0))&"","")'
        if src in ("E", "F"):
            f = f'=IFERROR(INDEX(Stock!${src}$2:${src}${N_STOCK + 1},MATCH($A{rr},{K},0)),"")'
        c = re_.cell(row=rr, column=j, value=f)
        c.font, c.border = F_FORMULE, BORDURE
    re_.cell(row=rr, column=2).font = Font(name="Consolas", size=11, bold=True)
    re_.cell(row=rr, column=6).number_format = "DD/MM/YYYY"
    re_.cell(row=rr, column=7).number_format = "#,##0"
re_.freeze_panes = "A7"
re_.conditional_formatting.add(f"H7:H{6 + N_RECH}", FormulaRule(formula=['ISNUMBER(FIND("PÉRIMÉ",H7))'], fill=FOND_ALERTE))
re_.conditional_formatting.add(f"H7:H{6 + N_RECH}", FormulaRule(formula=['ISNUMBER(FIND("Péremption",H7))'], fill=FOND_ATTENTION))

# ---------------------------------------------------------------- Étiquettes
et = wb.create_sheet("Étiquettes")
et.sheet_view.showGridLines = False
et.column_dimensions["A"].width = 30
et.column_dimensions["B"].width = 4
et.column_dimensions["C"].width = 30
et["A1"] = "Étiquettes d'emplacement à imprimer (2 par ligne) – générées depuis l'onglet Emplacements"
et["A1"].font = Font(name=POLICE, size=9, italic=True, color="808080")
N_ETIQ = 200
cadre = Side(style="medium", color="000000")
for k in range(N_ETIQ):
    rr = 3 + (k // 2) * 2
    cc = 1 if k % 2 == 0 else 3
    src = k + 2
    code = et.cell(row=rr, column=cc, value=f'=IF(Emplacements!$A${src}="","",Emplacements!$A${src})')
    code.font = Font(name=POLICE, size=28, bold=True)
    code.alignment = CENTRE
    code.border = Border(left=cadre, right=cadre, top=cadre)
    lib = et.cell(row=rr + 1, column=cc, value=f'=IF(Emplacements!$A${src}="","",Emplacements!$F${src})')
    lib.font = Font(name=POLICE, size=10)
    lib.alignment = CENTRE
    lib.border = Border(left=cadre, right=cadre, bottom=cadre)
    et.row_dimensions[rr].height = 44
    et.row_dimensions[rr + 1].height = 18
et.print_area = f"A1:C{3 + N_ETIQ + 1}"
et.page_setup.orientation = "portrait"
et.page_setup.fitToWidth = 1
et.page_setup.fitToHeight = 0
et.sheet_properties.pageSetUpPr.fitToPage = True

for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 100

wb.save(SORTIE)
print(f"Classeur écrit : {SORTIE}")
